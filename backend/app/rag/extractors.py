from pathlib import Path

from app.core.errors import AppError


def extract_file(path: Path, ext: str) -> list[dict]:
    ext = ext.lower()
    if ext in ("txt", "md"):
        text = path.read_text(encoding="utf-8", errors="ignore")
        return [{"text": text, "page_number": None, "slide_number": None}]
    if ext == "pdf":
        import fitz

        doc = fitz.open(path)
        pages = []
        for i, page in enumerate(doc, 1):
            pages.append({"text": page.get_text() or "", "page_number": i, "slide_number": None})
        doc.close()
        if not any(p["text"].strip() for p in pages):
            raise AppError(
                "EMPTY_PDF",
                "This PDF has no extractable text. It may be scanned; OCR is not supported in this prototype.",
            )
        return pages
    if ext == "docx":
        from docx import Document

        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [{"text": "\n".join(parts), "page_number": None, "slide_number": None}]
    if ext == "pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            bits = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    bits.append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        bits.append(" | ".join(cell.text for cell in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                bits.append(slide.notes_slide.notes_text_frame.text)
            slides.append({"text": "\n".join(bits), "page_number": None, "slide_number": i})
        return slides
    raise AppError("INVALID_FILE_TYPE", f"Unsupported type {ext}")
